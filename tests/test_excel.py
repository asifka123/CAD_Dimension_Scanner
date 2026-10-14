"""
tests/test_excel.py — Integration tests for the smart Excel updater.
Run with:  python -m pytest tests/
"""

import sys
import tempfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent.parent))

import openpyxl
import pytest
from dimscan.excel import update, META_ROW, LABEL_ROW, DATA_START, TYPE_COL


def _fake_matched(values, start_x=100, start_y=50):
    labels = [f"{chr(65+i*2)}-{chr(66+i*2)}" for i in range(len(values))]
    return [
        {"number": {"value": v, "x": start_x + i*50, "y": start_y},
         "line": None, "score": 0, "label": labels[i]}
        for i, v in enumerate(values)
    ]


def test_creates_new_file():
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    Path(path).unlink()  # ensure it doesn't exist

    matched = _fake_matched([10, 15, 20])
    update(matched, "drawing1.png", "Test Part NO. 1", excel_path=path)

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    assert ws.cell(row=LABEL_ROW, column=1).value == "Reading No"
    assert ws.cell(row=DATA_START, column=1).value == 1
    Path(path).unlink()


def test_same_identity_appends_row():
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    Path(path).unlink()

    update(_fake_matched([10, 15]), "img1.png", "Part A", excel_path=path)
    update(_fake_matched([12, 18]), "img2.png", "Part A", excel_path=path)

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    assert ws.cell(row=DATA_START,     column=1).value == 1
    assert ws.cell(row=DATA_START + 1, column=1).value == 2
    Path(path).unlink()


def test_different_identity_new_cols():
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    Path(path).unlink()

    update(_fake_matched([10, 15]), "img1.png", "Part A", excel_path=path)
    update(_fake_matched([50, 60]), "img2.png", "Part B", excel_path=path)

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    # Part B should have its identity in META_ROW at a different column than Part A
    meta_values = [ws.cell(row=META_ROW, column=c).value
                   for c in range(1, ws.max_column + 1)]
    assert "Part A" in meta_values
    assert "Part B" in meta_values
    assert meta_values.index("Part A") != meta_values.index("Part B")
    Path(path).unlink()


def test_same_image_updates_not_duplicates():
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    Path(path).unlink()

    update(_fake_matched([10, 15]), "img1.png", "Part A", excel_path=path)
    update(_fake_matched([11, 16]), "img1.png", "Part A", excel_path=path)  # re-run

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    # Should still be only ONE data row
    data_rows = [r for r in range(DATA_START, ws.max_row + 1)
                 if ws.cell(row=r, column=1).value is not None]
    assert len(data_rows) == 1
    Path(path).unlink()
