"""
excel.py — Smart Excel updater.

Layout
------
Row 1 (META, white text): drawing identity string per column group
Row 2 (LABELS): Reading No | Drawing Type | A-B | C-D | ... | (gap) | A-B | ...
Row 3+  (DATA): one row per image processed

Rules
-----
• Same identity → same column group, new row appended
• New identity  → new column group to the right
• Same filename → UPDATE existing row (idempotent re-runs)
"""

import openpyxl
from openpyxl.styles  import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils   import get_column_letter
from openpyxl.comments import Comment
from pathlib import Path

from . import config as cfg


# ── Style constants ───────────────────────────────────────────────────────────
_T   = Side(style="thin",   color="CCCCCC")
_TH  = Side(style="medium", color="444444")

HDR_FILL  = PatternFill("solid", fgColor="2F2F2F")
HDR_FONT  = Font(bold=True, color="FFFFFF", size=11)
META_FILL = PatternFill("solid", fgColor="FFFFFF")
META_FONT = Font(color="FFFFFF", size=8)       # invisible white text
TYPE_FILL = PatternFill("solid", fgColor="E8F4FD")
TYPE_FONT = Font(italic=True, color="1A5276",  size=9)
RNO_FILL  = PatternFill("solid", fgColor="F5F5F5")
RNO_FONT  = Font(bold=True, size=11)
ALT       = [PatternFill("solid", fgColor="FFFFFF"),
             PatternFill("solid", fgColor="FFF8F0")]
VF        = Font(size=11)
CA        = Alignment(horizontal="center", vertical="center")
LA        = Alignment(horizontal="left",   vertical="center")


def _b():  return Border(left=_T,  right=_T,  top=_T,  bottom=_T)
def _hb(): return Border(left=_TH, right=_TH, top=_TH, bottom=_TH)


# ── Row / column indices ──────────────────────────────────────────────────────
META_ROW  = 1     # hidden row: stores identity strings
LABEL_ROW = 2     # visible header row
DATA_START = 3    # first data row
TYPE_COL   = 2    # "Drawing Type" column
DIM_START  = 3    # first dimension column


# ── Internal helpers ──────────────────────────────────────────────────────────

def _find_group_by_identity(ws, identity):
    """Scan META_ROW for an exact identity match; return start column or None."""
    for col in range(DIM_START, ws.max_column + 1):
        cell = ws.cell(row=META_ROW, column=col)
        if cell.value and str(cell.value).strip() == identity:
            return col
    return None


def _count_group_cols(ws, start_col):
    """Count how many label columns exist in a group starting at start_col."""
    col = start_col
    while col <= ws.max_column and ws.cell(row=LABEL_ROW, column=col).value:
        col += 1
    return col - start_col


def _find_image_row(ws, img_name):
    """Return the data row that was created for img_name, or None."""
    for r in range(DATA_START, ws.max_row + 1):
        c = ws.cell(row=r, column=1)
        if c.comment and img_name in str(c.comment.text):
            return r
    return None


def _next_data_row(ws):
    r = DATA_START
    while ws.cell(row=r, column=1).value is not None:
        r += 1
    return r


def _next_free_dim_col(ws):
    mc = ws.max_column
    return max(DIM_START, mc + 2) if mc >= DIM_START else DIM_START


# ── Public API ────────────────────────────────────────────────────────────────

def update(matched, img_path, identity, excel_path=None):
    """
    Insert or update one row in the master Excel file.

    Parameters
    ----------
    matched     : list of {number, line, score, label}
    img_path    : str — source image path (used as row identifier)
    identity    : str — drawing identity (e.g. "Brake Caliper NO. 1")
    excel_path  : str or None — defaults to config.OUTPUT_EXCEL
    """
    excel_path = excel_path or cfg.OUTPUT_EXCEL
    img_name   = Path(img_path).name
    labels     = [m["label"] for m in matched]
    ep         = Path(excel_path)
    ep.parent.mkdir(parents=True, exist_ok=True)

    # ── Load or create workbook ───────────────────────────────────────────────
    if ep.exists():
        wb = openpyxl.load_workbook(str(ep))
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Dimensions"
        ws.freeze_panes = "C3"

        # Header: Reading No
        c = ws.cell(row=LABEL_ROW, column=1, value="Reading No")
        c.fill, c.font, c.border, c.alignment = HDR_FILL, HDR_FONT, _hb(), CA
        ws.column_dimensions["A"].width = 14

        # Header: Drawing Type
        c = ws.cell(row=LABEL_ROW, column=TYPE_COL, value="Drawing Type")
        c.fill, c.font, c.border, c.alignment = HDR_FILL, HDR_FONT, _hb(), CA
        ws.column_dimensions["B"].width = 26

        print(f"  [excel] created: {ep.name}")

    # ── Find or create column group ───────────────────────────────────────────
    sc = _find_group_by_identity(ws, identity)

    if sc is None:
        # Brand-new drawing type → new column group
        sc = _next_free_dim_col(ws)

        # Write identity into META_ROW (invisible)
        mc = ws.cell(row=META_ROW, column=sc, value=identity)
        mc.fill, mc.font = META_FILL, META_FONT

        # Write labels into LABEL_ROW
        for i, lbl in enumerate(labels):
            c = ws.cell(row=LABEL_ROW, column=sc + i, value=lbl)
            c.fill, c.font, c.border, c.alignment = HDR_FILL, HDR_FONT, _hb(), CA
            ws.column_dimensions[get_column_letter(sc + i)].width = 10

        print(f"  [excel] new group '{identity}' "
              f"@ cols {get_column_letter(sc)}–{get_column_letter(sc+len(labels)-1)}")

    else:
        # Known type — extend columns if this image has more dims
        existing = _count_group_cols(ws, sc)
        if len(labels) > existing:
            for i in range(existing, len(labels)):
                c = ws.cell(row=LABEL_ROW, column=sc + i, value=labels[i])
                c.fill, c.font, c.border, c.alignment = HDR_FILL, HDR_FONT, _hb(), CA
                ws.column_dimensions[get_column_letter(sc + i)].width = 10
            print(f"  [excel] extended '{identity}' to {len(labels)} cols")
        else:
            print(f"  [excel] known group '{identity}'")

    # ── Find or create data row ───────────────────────────────────────────────
    dr = _find_image_row(ws, img_name)
    if dr:
        print(f"  [excel] same image — updating row {dr}")
    else:
        dr  = _next_data_row(ws)
        rno = dr - DATA_START + 1

        rn = ws.cell(row=dr, column=1, value=rno)
        rn.fill, rn.font, rn.border, rn.alignment = RNO_FILL, RNO_FONT, _b(), CA
        rn.comment = Comment(img_name, "dimscan")
        print(f"  [excel] new row — Reading No {rno}")

    # Drawing type label
    alt = ALT[(dr - DATA_START) % 2]
    tb  = ws.cell(row=dr, column=TYPE_COL, value=identity[:32])
    tb.fill, tb.font, tb.border, tb.alignment = TYPE_FILL, TYPE_FONT, _b(), LA

    # Write dimension values
    for i, m in enumerate(matched):
        vc = ws.cell(row=dr, column=sc + i, value=m["number"]["value"])
        vc.fill, vc.font, vc.border, vc.alignment = alt, VF, _b(), CA

    wb.save(str(ep))
    print(f"  [excel] saved → {ep.name}")
